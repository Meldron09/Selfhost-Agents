"""The `output-writer` subagent: four per-format write tools.

Resolves wayfinder ticket #5's confirmed design (see the prototype at
`prototypes/output-writer.PROTOTYPE.html` on
`prototype/output-writer-subagent-design`):

* A caller passes a bare `name` with no extension; each tool appends its own
  canonical one (`.xlsx`/`.docx`/`.pptx`/`.txt`) — removing an
  extension-mismatch failure mode entirely rather than needing to validate one.
* Structured input shapes mirror `file-reader`'s *return* shapes (a later
  ticket, #17), so a round trip (read a file, write a modified version) needs
  no reshaping.
* Every tool validates its structured input *before* storing any bytes:
  malformed shape or empty content returns an explicit error, no key issued,
  `outputs` untouched. Only on success does it call `store_output_bytes` and
  append `{key, filename}` to the `outputs` state field (docs/adr/0001) via a
  `Command` update — the standard LangGraph pattern for a tool that both
  replies and mutates state, since `output-writer` runs as a declarative
  `SubAgent` and its `outputs` writes flow back to the orchestrator's own
  state through the `task` tool's own state-forwarding (confirmed in
  docs/research/deepagents-state-schema.md).
"""
from __future__ import annotations

import io
from typing import Any, NotRequired, TypedDict

from deepagents.middleware.subagents import SubAgent
from docx import Document
from langchain.tools import ToolRuntime
from langchain_core.messages import ToolMessage
from langchain_core.tools import tool
from langgraph.types import Command
from openpyxl import Workbook
from pptx import Presentation

from agent.files.store import store_output_bytes


class SheetInput(TypedDict):
    sheet_name: str
    headers: list[str]
    rows: NotRequired[list[list[Any]]]


class SlideInput(TypedDict):
    title: str
    bullets: list[str]
    speaker_notes: NotRequired[str]


# --- validation (before any bytes are stored) -----------------------------


def _validate_xlsx(sheets: list[SheetInput]) -> str | None:
    if not sheets:
        return "sheets must be a non-empty list"
    for sheet in sheets:
        name = sheet.get("sheet_name")
        if not name:
            return "every sheet needs a sheet_name"
        headers = sheet.get("headers")
        if not headers:
            return f'sheet "{name}" needs non-empty headers'
        for row in sheet.get("rows") or []:
            if len(row) != len(headers):
                return f'sheet "{name}": a row has {len(row)} values but {len(headers)} headers'
    return None


def _validate_text(text: Any) -> str | None:
    if not isinstance(text, str) or not text.strip():
        return "text must be a non-empty string"
    return None


def _validate_pptx(slides: list[SlideInput]) -> str | None:
    if not slides:
        return "slides must be a non-empty list"
    for slide in slides:
        title = slide.get("title")
        if not title:
            return "every slide needs a title"
        if not isinstance(slide.get("bullets"), list):
            return f'slide "{title}" needs a bullets list (can be empty, not missing)'
    return None


# --- serialization (only reached once validation has passed) --------------


def _serialize_xlsx(sheets: list[SheetInput]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet["sheet_name"])
        worksheet.append(sheet["headers"])
        for row in sheet.get("rows") or []:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _serialize_docx(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _serialize_pptx(slides: list[SlideInput]) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title + content
    for slide_input in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = slide_input["title"]
        bullets = slide_input.get("bullets") or []
        if bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet
        notes = slide_input.get("speaker_notes") or ""
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# --- shared success/failure reply, mirroring the prototype's wording ------


def _write_result(
    tool_name: str,
    name: str,
    filename: str,
    error: str | None,
    data: bytes | None,
    runtime: ToolRuntime,
) -> Command:
    if error is not None:
        return Command(
            update={
                "messages": [
                    ToolMessage(
                        content=f'Could not write "{name}" via {tool_name}: {error}',
                        name=tool_name,
                        tool_call_id=runtime.tool_call_id,
                        status="error",
                    )
                ],
            }
        )
    key = store_output_bytes(filename, data)
    return Command(
        update={
            "outputs": [{"key": key, "filename": filename}],
            "messages": [
                ToolMessage(
                    content=f'Wrote "{filename}" via {tool_name} and registered it as an Output (key {key}).',
                    name=tool_name,
                    tool_call_id=runtime.tool_call_id,
                    status="success",
                )
            ],
        }
    )


# --- the four tools ---------------------------------------------------------


@tool
def write_xlsx(name: str, sheets: list[SheetInput], runtime: ToolRuntime) -> Command:
    """Write tabular data as an .xlsx workbook and register it as an Output.

    `sheets` is a non-empty list of `{sheet_name, headers, rows}` — one entry
    per worksheet. Every row must have exactly as many values as `headers`.
    `name` carries no extension; `.xlsx` is appended automatically.
    """
    error = _validate_xlsx(sheets)
    data = None if error else _serialize_xlsx(sheets)
    return _write_result("write_xlsx", name, f"{name}.xlsx", error, data, runtime)


@tool
def write_docx(name: str, text: str, runtime: ToolRuntime) -> Command:
    """Write prose/markdown-ish text as a .docx document and register it as an Output.

    `name` carries no extension; `.docx` is appended automatically.
    """
    error = _validate_text(text)
    data = None if error else _serialize_docx(text)
    return _write_result("write_docx", name, f"{name}.docx", error, data, runtime)


@tool
def write_pptx(name: str, slides: list[SlideInput], runtime: ToolRuntime) -> Command:
    """Write a slide deck as a .pptx presentation and register it as an Output.

    `slides` is a non-empty list of `{title, bullets, speaker_notes}` — one
    entry per slide (`bullets` may be empty, but must be present).
    `name` carries no extension; `.pptx` is appended automatically.
    """
    error = _validate_pptx(slides)
    data = None if error else _serialize_pptx(slides)
    return _write_result("write_pptx", name, f"{name}.pptx", error, data, runtime)


@tool
def write_txt(name: str, text: str, runtime: ToolRuntime) -> Command:
    """Write plain text as a .txt file and register it as an Output.

    `name` carries no extension; `.txt` is appended automatically.
    """
    error = _validate_text(text)
    data = None if error else text.encode("utf-8")
    return _write_result("write_txt", name, f"{name}.txt", error, data, runtime)


SYSTEM_PROMPT = """You are the output-writer subagent. You are given structured data and you turn it into a downloadable file for the orchestrator, registering it as an Output.

Method:
1. Choose the tool by the file's format: write_xlsx (tabular data, list of sheets), write_docx (prose/markdown-ish text), write_pptx (a slide deck), write_txt (plain text).
2. Pass the structured data straight through in the shape each tool expects -- do not flatten a table to text for write_xlsx, or split prose into fake slides for write_pptx.
3. Give each tool a short descriptive name with no file extension; the tool appends its own extension and stores the bytes -- you never construct a path or a key yourself.
4. If a tool reports it could not write the file (malformed input, e.g. a spreadsheet row that doesn't match its headers, or empty content), report that by name and the reason given. Do not retry with guessed fixes and do not report the file as written if it wasn't.

Return a synthesis, not a raw dump. Your final message must contain, for every file you attempted:
- The requested file and which tool handled it.
- On success: the filename actually written and confirmation it was registered as an Output, so the orchestrator can tell the person it's ready to download.
- On failure: an explicit line naming the attempted file and the reason, so the orchestrator can relay it rather than silently claiming success.

Never report a file as written when the underlying write failed. A gap that is labelled is useful; a claimed Output that doesn't exist wastes the person's next click."""


def build_output_writer_subagent() -> SubAgent:
    """The `SubAgent` spec the orchestrator's `task` tool delegates to."""
    return {
        "name": "output-writer",
        "description": (
            "Produce a downloadable file from structured content, registered as an "
            "Output. Always delegate any file here — the orchestrator has no "
            "other way to write a file."
        ),
        "system_prompt": SYSTEM_PROMPT,
        "tools": [write_xlsx, write_docx, write_pptx, write_txt],
    }
