"""The `file-reader` subagent: five per-format read tools.

Resolves issue #17's confirmed design, mirroring `agent/output_writer.py`'s
shape and split — pure validation/extraction functions the tool-level
wrappers delegate to, so tests/test_file_reader.py can exercise
success/error cases without a `ToolRuntime`.

* A caller passes only an opaque `key` (never a filename or path) — the
  attachment's filename lives in the orchestrator's own `task` description,
  not in a tool argument, matching CONTEXT.md's "Key" definition: callers
  never construct or parse one, only pass it back.
* Dispatch by filename extension is a prompt-level decision, made by the
  file-reader subagent itself (SYSTEM_PROMPT below) — an unsupported
  extension short-circuits before any tool call, so there is no
  `read_unsupported` tool and no extension parameter to validate here.
* Return shapes mirror `output-writer`'s structured *input* shapes
  (`SheetInput`/`SlideInput` in agent/output_writer.py), so a round trip
  (read a file, hand its content straight to `write_xlsx`/`write_pptx`)
  needs no reshaping — `SheetOutput`/`SlideOutput` below are supersets that
  add read-only fields (`slide_number`) those write-side shapes don't need.
* Every tool resolves the key and parses the bytes *before* replying;
  either failure (unknown key, or a supported extension whose bytes don't
  parse) becomes an explicit tool-level error via a `Command`-wrapped
  `ToolMessage(status="error")` — never a raised exception propagating into
  the model loop, and never a partial/fabricated result standing in for
  content that wasn't actually extracted.
* Every `read_*` tool is a one-line call into the shared `_read` funnel,
  rather than inlining its own resolve/extract/reply sequence the way
  `output_writer.py`'s `write_*` tools do. That's a deliberate divergence
  from the otherwise-mirrored shape: `output-writer`'s four tools each have
  a genuinely different validate/serialize step, while all five `read_*`
  tools differ only in *which* extractor they call, so funnelling that one
  varying piece through `_read` removes five copies of identical
  resolve-then-reply plumbing without losing any of it.
"""
from __future__ import annotations

import io
import json
from collections.abc import Callable
from typing import Any, TypedDict

import xlrd
from deepagents.middleware.subagents import SubAgent
from docx import Document
from langchain.tools import ToolRuntime
from langchain_core.messages import ToolMessage
from langchain_core.tools import tool
from langgraph.types import Command
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from agent.files.store import resolve_attachment_bytes


class SheetOutput(TypedDict):
    sheet_name: str
    headers: list[str]
    rows: list[list[Any]]


class SlideOutput(TypedDict):
    slide_number: int
    title: str
    bullets: list[str]
    speaker_notes: str


# --- key resolution (shared by every tool) ---------------------------------


def _resolve(key: str) -> tuple[bytes | None, str | None]:
    try:
        return resolve_attachment_bytes(key), None
    except KeyError:
        return None, f'no attachment found for key "{key}"'


# --- extraction (only reached once the key has resolved) -------------------


def _safely(compute: Callable[[], Any]) -> tuple[Any | None, str | None]:
    """Run `compute`, turning any parse failure into `(None, str(exc))`
    rather than letting it raise — the one shape every `_extract_*` function
    below shares, factored out so each states only what's actually different
    about its own format.
    """
    try:
        return compute(), None
    except Exception as exc:  # noqa: BLE001 - any parse failure becomes an explicit error, never a raised exception
        return None, str(exc)


def _extract_pdf(data: bytes) -> tuple[str | None, str | None]:
    def compute() -> str:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    return _safely(compute)


def _extract_xlsx_ooxml(data: bytes) -> tuple[list[SheetOutput] | None, str | None]:
    """.xlsx/.xlsm, via `openpyxl` — the modern, zip-based Excel format."""

    def compute() -> list[SheetOutput]:
        workbook = load_workbook(io.BytesIO(data), data_only=True)
        sheets: list[SheetOutput] = []
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            headers = [str(cell) if cell is not None else "" for cell in (rows[0] if rows else [])]
            body = [list(row) for row in rows[1:]]
            sheets.append({"sheet_name": worksheet.title, "headers": headers, "rows": body})
        return sheets

    return _safely(compute)


def _extract_xls_legacy(data: bytes) -> tuple[list[SheetOutput] | None, str | None]:
    """.xls, via `xlrd` — the legacy binary (OLE2/BIFF) Excel format
    `openpyxl` cannot read at all (it only understands the zip-based
    xlsx/xlsm formats). `xlrd` 2.x deliberately dropped xlsx support in the
    other direction, so the two libraries are complementary, not overlapping.
    """

    def compute() -> list[SheetOutput]:
        book = xlrd.open_workbook(file_contents=data)
        sheets: list[SheetOutput] = []
        for sheet in book.sheets():
            rows = [sheet.row_values(i) for i in range(sheet.nrows)]
            headers = [str(v) for v in (rows[0] if rows else [])]
            body = [list(row) for row in rows[1:]]
            sheets.append({"sheet_name": sheet.name, "headers": headers, "rows": body})
        return sheets

    return _safely(compute)


def _extract_xlsx(data: bytes) -> tuple[list[SheetOutput] | None, str | None]:
    """`read_xlsx` handles both `.xlsx` and `.xls` (issue #17's acceptance
    criteria), which are different binary formats needing different
    libraries — try the modern one first, since it's what the extension
    itself names, and fall back to the legacy reader only if that fails.
    """
    sheets, ooxml_error = _extract_xlsx_ooxml(data)
    if sheets is not None:
        return sheets, None
    legacy_sheets, legacy_error = _extract_xls_legacy(data)
    if legacy_sheets is not None:
        return legacy_sheets, None
    return None, f"not a readable .xlsx ({ooxml_error}) or .xls ({legacy_error})"


def _extract_docx(data: bytes) -> tuple[str | None, str | None]:
    def compute() -> str:
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    return _safely(compute)


def _extract_pptx(data: bytes) -> tuple[list[SlideOutput] | None, str | None]:
    def compute() -> list[SlideOutput]:
        presentation = Presentation(io.BytesIO(data))
        slides: list[SlideOutput] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title = slide.shapes.title.text if slide.shapes.title else ""
            bullets: list[str] = []
            for shape in slide.shapes:
                if shape == slide.shapes.title or not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text
                    if text:
                        bullets.append(text)
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
            slides.append({"slide_number": index, "title": title, "bullets": bullets, "speaker_notes": notes})
        return slides

    return _safely(compute)


def _extract_txt(data: bytes) -> tuple[str | None, str | None]:
    return _safely(lambda: data.decode("utf-8"))


# --- shared success/failure reply, mirroring output_writer.py's shape ------


def _read_result(
    tool_name: str,
    key: str,
    error: str | None,
    content: Any,
    runtime: ToolRuntime,
) -> Command:
    if error is not None:
        return Command(
            update={
                "messages": [
                    ToolMessage(
                        content=f'Could not read attachment (key "{key}") via {tool_name}: {error}',
                        name=tool_name,
                        tool_call_id=runtime.tool_call_id,
                        status="error",
                    )
                ],
            }
        )
    body = content if isinstance(content, str) else json.dumps(content)
    return Command(
        update={
            "messages": [
                ToolMessage(
                    content=body,
                    name=tool_name,
                    tool_call_id=runtime.tool_call_id,
                    status="success",
                )
            ],
        }
    )


def _read(tool_name: str, key: str, runtime: ToolRuntime, extractor) -> Command:
    data, resolve_error = _resolve(key)
    if resolve_error is not None:
        return _read_result(tool_name, key, resolve_error, None, runtime)
    content, extract_error = extractor(data)
    return _read_result(tool_name, key, extract_error, content, runtime)


# --- the five tools ----------------------------------------------------------


@tool
def read_pdf(key: str, runtime: ToolRuntime) -> Command:
    """Extract text from a PDF Attachment.

    `key` is the attachment's opaque Key, given to you as part of the
    orchestrator's task description — never a filename or path.
    """
    return _read("read_pdf", key, runtime, _extract_pdf)


@tool
def read_xlsx(key: str, runtime: ToolRuntime) -> Command:
    """Extract tabular data from an .xlsx/.xls Attachment.

    Returns a list of `{sheet_name, headers, rows}` — one entry per
    worksheet, `headers` being the first row and `rows` every row after it.
    """
    return _read("read_xlsx", key, runtime, _extract_xlsx)


@tool
def read_docx(key: str, runtime: ToolRuntime) -> Command:
    """Extract text from a .docx Attachment, one paragraph per line."""
    return _read("read_docx", key, runtime, _extract_docx)


@tool
def read_pptx(key: str, runtime: ToolRuntime) -> Command:
    """Extract slide content from a .pptx Attachment.

    Returns a list of `{slide_number, title, bullets, speaker_notes}` — one
    entry per slide.
    """
    return _read("read_pptx", key, runtime, _extract_pptx)


@tool
def read_txt(key: str, runtime: ToolRuntime) -> Command:
    """Extract the content of a .txt/.md Attachment, decoded as UTF-8 text."""
    return _read("read_txt", key, runtime, _extract_txt)


SYSTEM_PROMPT = """You are the file-reader subagent. You are given an Attachment (a filename and an opaque key, in the task description) and you extract its actual content so the orchestrator can answer a question grounded in it.

Method:
1. Look at the filename's extension to choose the right tool: read_pdf (.pdf), read_xlsx (.xlsx or .xls), read_docx (.docx), read_pptx (.pptx), read_txt (.txt or .md).
2. If the extension matches none of those, do not call any tool. Reply directly that the file type is unsupported, naming the exact filename.
3. Call the matching tool with only the key you were given -- never guess, construct, or reuse a key from a different attachment.
4. If a tool reports it could not extract content (an unknown key, or a file whose bytes don't parse), report that by filename and the reason given. Do not retry with guessed fixes, and never invent or approximate content that wasn't actually extracted.

Return a synthesis grounded in what the tool actually returned, not a raw dump and not a guess. Your final message must, for every Attachment you were asked to read:
- Name the file and which tool handled it (or that its type is unsupported).
- On success: answer using the actual extracted content, or if no question was asked, summarize what the file actually contains.
- On failure: an explicit line naming the attempted file and the reason, so the orchestrator can relay it rather than silently claiming an answer.

Never present an answer as grounded in a file's content when extraction failed or was never attempted. A gap that is labelled is useful; a fabricated answer wastes the person's trust."""


def build_file_reader_subagent() -> SubAgent:
    """The `SubAgent` spec the orchestrator's `task` tool delegates to."""
    return {
        "name": "file-reader",
        "description": (
            "Extract the actual content of an uploaded Attachment (pdf, xlsx/xls, "
            "docx, pptx, or txt/md) so a question about it can be answered from "
            "real content. Always delegate every Attachment here -- the "
            "orchestrator has no other way to read a file."
        ),
        "system_prompt": SYSTEM_PROMPT,
        "tools": [read_pdf, read_xlsx, read_docx, read_pptx, read_txt],
    }
