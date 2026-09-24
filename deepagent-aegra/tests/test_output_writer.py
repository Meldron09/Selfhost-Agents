"""`agent.output_writer`'s pure validation/serialization functions.

No `ToolRuntime`/graph involved here — these are the seams the four
`@tool`-wrapped functions delegate to before touching the file store. The
tool-level, state-updating behavior (outputs accumulation, no key on
failure) is exercised end to end in tests/test_graph.py, driven through
`build_agent` with a `ScriptedChatModel` — the style
docs/adr/0005-testing-strategy-carryover-from-agent-runtime.md calls for.
"""
from __future__ import annotations

import io

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from agent.output_writer import (
    _serialize_docx,
    _serialize_pptx,
    _serialize_xlsx,
    _validate_pptx,
    _validate_text,
    _validate_xlsx,
)

# --- write_xlsx ---------------------------------------------------------


def test_validate_xlsx_accepts_well_formed_sheets():
    sheets = [
        {"sheet_name": "Revenue", "headers": ["Region", "Q2"], "rows": [["EMEA", 100]]},
    ]
    assert _validate_xlsx(sheets) is None


def test_validate_xlsx_rejects_an_empty_sheet_list():
    assert "non-empty list" in _validate_xlsx([])


def test_validate_xlsx_rejects_a_sheet_with_no_name():
    sheets = [{"sheet_name": "", "headers": ["A"], "rows": []}]
    assert "sheet_name" in _validate_xlsx(sheets)


def test_validate_xlsx_rejects_empty_headers():
    sheets = [{"sheet_name": "Revenue", "headers": [], "rows": []}]
    assert "headers" in _validate_xlsx(sheets)


def test_validate_xlsx_rejects_a_row_headers_length_mismatch():
    sheets = [
        {"sheet_name": "Revenue", "headers": ["Region", "Q2", "Q3"], "rows": [["EMEA", 1450000]]},
    ]
    error = _validate_xlsx(sheets)
    assert "Revenue" in error
    assert "2" in error and "3" in error


def test_serialize_xlsx_produces_a_real_workbook_openable_by_openpyxl():
    sheets = [
        {"sheet_name": "Revenue", "headers": ["Region", "Q2"], "rows": [["EMEA", 100], ["AMER", 200]]},
        {"sheet_name": "Headcount", "headers": ["Team", "Count"], "rows": [["Eng", 42]]},
    ]
    data = _serialize_xlsx(sheets)

    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Revenue", "Headcount"]
    revenue = workbook["Revenue"]
    assert [cell.value for cell in revenue[1]] == ["Region", "Q2"]
    assert [cell.value for cell in revenue[2]] == ["EMEA", 100]
    assert [cell.value for cell in revenue[3]] == ["AMER", 200]


# --- write_docx / write_txt (shared text validation) ---------------------


def test_validate_text_rejects_empty_string():
    assert _validate_text("") is not None


def test_validate_text_rejects_whitespace_only():
    assert _validate_text("   \n\t") is not None


def test_validate_text_accepts_non_empty_string():
    assert _validate_text("hello") is None


def test_serialize_docx_produces_a_real_document_openable_by_python_docx():
    data = _serialize_docx("# Kickoff Memo\n\nScope: independent demo.")

    document = Document(io.BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs]
    assert "# Kickoff Memo" in paragraphs
    assert "Scope: independent demo." in paragraphs


# --- write_pptx -----------------------------------------------------------


def test_validate_pptx_accepts_well_formed_slides():
    slides = [{"title": "Intro", "bullets": ["Point one"], "speaker_notes": ""}]
    assert _validate_pptx(slides) is None


def test_validate_pptx_rejects_an_empty_slide_list():
    assert "non-empty list" in _validate_pptx([])


def test_validate_pptx_rejects_a_slide_with_no_title():
    slides = [{"title": "", "bullets": []}]
    assert "title" in _validate_pptx(slides)


def test_validate_pptx_rejects_a_missing_bullets_list():
    slides = [{"title": "Intro"}]
    error = _validate_pptx(slides)
    assert "bullets" in error


def test_validate_pptx_accepts_empty_bullets():
    slides = [{"title": "Intro", "bullets": []}]
    assert _validate_pptx(slides) is None


def test_serialize_pptx_produces_a_real_presentation_openable_by_python_pptx():
    slides = [
        {"title": "Deepagent on Aegra", "bullets": ["Independent demo", "No sandbox"], "speaker_notes": ""},
        {"title": "Subagents", "bullets": ["output-writer"], "speaker_notes": "mirrors deep-research"},
    ]
    data = _serialize_pptx(slides)

    presentation = Presentation(io.BytesIO(data))
    assert len(presentation.slides) == 2
    first = presentation.slides[0]
    assert first.shapes.title.text == "Deepagent on Aegra"
    second = presentation.slides[1]
    assert second.notes_slide.notes_text_frame.text == "mirrors deep-research"
